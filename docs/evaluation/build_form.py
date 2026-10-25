"""Build the Kobo XLSForm: evaluation of CAMLEAP autonomous entity registration."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT = sys.argv[1]
FR, EN = "label::Français (fr)", "label::English (en)"
HFR, HEN = "hint::Français (fr)", "hint::English (en)"
CFR, CEN = "constraint_message::Français (fr)", "constraint_message::English (en)"
COLS = ["type", "name", FR, EN, HFR, HEN, "required", "relevant", "constraint",
        CFR, CEN, "appearance", "choice_filter", "calculation", "default", "parameters"]

survey, choices = [], []


def q(type_, name, fr="", en="", hint=None, req=None, **kw):
    row = {"type": type_, "name": name, FR: fr, EN: en or fr}
    if hint:
        row[HFR], row[HEN] = hint
    if req is None:
        req = type_.split()[0] in ("select_one", "select_multiple", "integer", "text") and type_ != "text"
    if req:
        row["required"] = "yes"
    for k, v in kw.items():
        key = {"cmsg": None}.get(k, k)
        if k == "cmsg":
            row[CFR], row[CEN] = v
        else:
            row[k] = v
    survey.append(row)


def begin(name, fr, en, appearance="field-list", relevant=None):
    q("begin_group", name, fr, en, req=False, appearance=appearance, **({"relevant": relevant} if relevant else {}))


def end(name):
    survey.append({"type": "end_group", "name": name})


def note(name, fr, en, **kw):
    q("note", name, fr, en, req=False, **kw)


def lst(name, items, entities=None):
    for it in items:
        code, fr, en = it[:3]
        row = {"list_name": name, "name": code, FR: fr, EN: en}
        if len(it) > 3:
            row["entities"] = it[3]
        choices.append(row)


# ---------------------------------------------------------------- choices
ENT = [
    ("ent", "Entreprise", "Company"),
    ("coop", "Coopérative", "Cooperative"),
    ("ong", "ONG ou association", "NGO or association"),
    ("adm", "Administration publique", "Public administration"),
    ("ctd", "Collectivité territoriale (CTD)", "Local authority (RLA)"),
    ("pp", "Projet, programme ou structure sous-tutelle", "Project, programme or supervised structure"),
    ("cfp", "Centre de formation professionnelle (CFP)", "Vocational training centre (VTC)"),
]
lst("entity", ENT)
lst("yesno", [("oui", "Oui", "Yes"), ("non", "Non", "No")])
lst("browser", [("chrome", "Google Chrome", "Google Chrome"), ("edge", "Microsoft Edge", "Microsoft Edge"),
                ("firefox", "Mozilla Firefox", "Mozilla Firefox"), ("opera", "Opera", "Opera"),
                ("autre", "Autre", "Other")])
lst("screen", [("small", "Petit écran (portable 13-14 pouces)", "Small (13-14 inch laptop)"),
               ("medium", "Écran moyen (15-17 pouces)", "Medium (15-17 inch)"),
               ("large", "Grand écran (moniteur 19 pouces et plus)", "Large (19 inch monitor or more)")])
lst("network", [("fibre", "Réseau du bureau (fibre / ADSL / Wi-Fi)", "Office network (fibre / ADSL / Wi-Fi)"),
                ("mobile", "Partage de connexion / clé 3G-4G", "Mobile hotspot / 3G-4G dongle"),
                ("autre", "Autre", "Other")])
lst("lang", [("fr", "Français", "French"), ("en", "Anglais", "English"),
             ("both", "Les deux (j'ai changé de langue)", "Both (I switched language)")])
lst("outcome", [("full", "Oui, jusqu'au récépissé final", "Yes, up to the final receipt"),
                ("partial", "Non, j'ai été bloqué(e) en cours de route", "No, I was blocked along the way"),
                ("abandon", "Non, j'ai abandonné", "No, I gave up")])
STEPS = [
    ("s1", "Étape 1 — Type d'organisation", "Step 1 — Type of organisation"),
    ("s2", "Étape 2 — Informations du répondant", "Step 2 — Respondent details"),
    ("s3", "Étape 3 — Organisation (identification)", "Step 3 — Organisation (identification)"),
    ("s4", "Étape 4 — Localisation", "Step 4 — Location"),
    ("s5", "Étape 5 — Sécurisez votre compte (mot de passe)", "Step 5 — Secure your account (password)"),
    ("s6", "Étape 6 — Récapitulatif", "Step 6 — Summary"),
    ("s7", "Récépissé final (« Et ensuite »)", "Final receipt (\"What next\")"),
]
lst("step", STEPS)
lst("step_general", STEPS + [("gen", "Ensemble du parcours / page d'accueil", "Whole journey / home page"),
                              ("login", "Première connexion après inscription", "First sign-in after registration")])
lst("ease", [("1", "Très difficile", "Very difficult"), ("2", "Difficile", "Difficult"),
             ("3", "Moyen", "Average"), ("4", "Facile", "Easy"), ("5", "Très facile", "Very easy"),
             ("na", "Étape non atteinte", "Step not reached")])
lst("agree", [("1", "Pas du tout d'accord", "Strongly disagree"), ("2", "Pas d'accord", "Disagree"),
              ("3", "Neutre", "Neutral"), ("4", "D'accord", "Agree"), ("5", "Tout à fait d'accord", "Strongly agree"),
              ("na", "Sans objet", "Not applicable")])
lst("sus", [("1", "1 — Pas du tout d'accord", "1 — Strongly disagree"), ("2", "2", "2"), ("3", "3", "3"),
            ("4", "4", "4"), ("5", "5 — Tout à fait d'accord", "5 — Strongly agree")])
lst("seq", [(str(i), str(i) + (" — Très difficile" if i == 1 else " — Très facile" if i == 7 else ""),
             str(i) + (" — Very difficult" if i == 1 else " — Very easy" if i == 7 else "")) for i in range(1, 8)])
lst("score5", [("1", "1 — Très insuffisant", "1 — Very poor"), ("2", "2 — Insuffisant", "2 — Poor"),
               ("3", "3 — Acceptable", "3 — Acceptable"), ("4", "4 — Bon", "4 — Good"),
               ("5", "5 — Excellent", "5 — Excellent")])
lst("speed", [("1", "Très lent", "Very slow"), ("2", "Lent", "Slow"), ("3", "Acceptable", "Acceptable"),
              ("4", "Rapide", "Fast"), ("5", "Très rapide", "Very fast")])
lst("submit_wait", [("lt5", "Moins de 5 secondes", "Under 5 seconds"), ("5_15", "5 à 15 secondes", "5 to 15 seconds"),
                    ("15_30", "15 à 30 secondes", "15 to 30 seconds"), ("gt30", "Plus de 30 secondes", "Over 30 seconds"),
                    ("never", "La soumission n'a pas abouti", "Submission never completed")])
lst("incident", [
    ("none", "Aucun incident", "No incident"),
    ("slow", "Page très lente à charger", "Page very slow to load"),
    ("freeze", "Page figée ou blanche", "Page frozen or blank"),
    ("err_unclear", "Message d'erreur incompréhensible", "Error message I could not understand"),
    ("valid_refused", "Une valeur correcte a été refusée", "A correct value was refused"),
    ("invalid_accepted", "Une valeur manifestement fausse a été acceptée", "An obviously wrong value was accepted"),
    ("data_lost", "Des informations saisies ont été perdues", "Data I had entered was lost"),
    ("list_empty", "Liste déroulante vide ou qui ne se charge pas (région, département…)", "Drop-down empty or not loading (region, division…)"),
    ("dup", "E-mail ou téléphone signalé comme déjà utilisé", "E-mail or phone reported as already in use"),
    ("display", "Affichage cassé (texte coupé, éléments superposés)", "Broken layout (cut text, overlapping elements)"),
    ("lang_mix", "Mélange de français et d'anglais", "French and English mixed"),
    ("tech_code", "Code technique ou texte incompréhensible affiché", "Technical code or meaningless text shown"),
    ("other", "Autre", "Other"),
])
lst("help_src", [("link", "Le lien d'aide de la plateforme", "The platform's help link"),
                 ("facilitator", "Le facilitateur du test", "The test facilitator"),
                 ("colleague", "Un collègue", "A colleague"), ("other", "Autre", "Other")])
lst("email_rx", [("yes", "Oui", "Yes"), ("no", "Non", "No"), ("unchecked", "Je n'ai pas vérifié", "I did not check")])
lst("realism", [("yes", "Oui, sans difficulté", "Yes, easily"),
                ("search", "Oui, mais il faudrait chercher des documents", "Yes, but would need to look up documents"),
                ("no", "Non, pas toujours", "No, not always")])
lst("prop_cat", [
    ("content", "Information demandée (à ajouter, retirer, rendre facultative…)", "Information requested (add, remove, make optional…)"),
    ("wording", "Libellé, terminologie, traduction", "Wording, terminology, translation"),
    ("options", "Liste de choix incomplète ou inadaptée", "Incomplete or unsuitable option list"),
    ("navigation", "Navigation, enchaînement des étapes", "Navigation, step sequence"),
    ("visual", "Présentation visuelle, lisibilité", "Visual layout, readability"),
    ("errors", "Messages d'erreur, contrôles de saisie", "Error messages, input checks"),
    ("bug", "Dysfonctionnement (bug)", "Malfunction (bug)"),
    ("perf", "Performance, lenteur", "Performance, slowness"),
    ("access", "Accessibilité (clavier, contraste, taille du texte)", "Accessibility (keyboard, contrast, text size)"),
    ("other", "Autre", "Other"),
])
lst("priority", [("p1", "Bloquant — empêche l'inscription", "Blocking — prevents registration"),
                 ("p2", "Important — à corriger avant le pilote", "Important — fix before the pilot"),
                 ("p3", "Souhaitable — amélioration", "Desirable — improvement"),
                 ("p4", "Mineur — détail", "Minor — detail")])
lst("ready", [("yes", "Oui, en l'état", "Yes, as it is"),
              ("minor", "Oui, après des corrections mineures", "Yes, after minor corrections"),
              ("major", "Non, des corrections importantes sont nécessaires", "No, major corrections are needed"),
              ("no", "Non, le parcours doit être revu", "No, the journey must be redesigned")])

# The fields of the registration form (react-web/src/lib/register-constants.ts
# and the respondent / location / security steps of /register), tagged with the
# entity types that see them; choice_filter shows each tester only theirs.
ALL = "all"
FIELDS = [
    # Step 2 — respondent
    ("r_prenom", "Répondant — Prénom", "Respondent — First name", ALL),
    ("r_nom", "Répondant — Nom", "Respondent — Last name", ALL),
    ("r_fonction", "Répondant — Fonction", "Respondent — Position", ALL),
    ("r_email", "Répondant — Email professionnel", "Respondent — Work e-mail", ALL),
    ("r_tel1", "Répondant — Téléphone", "Respondent — Phone", ALL),
    ("r_tel2", "Répondant — Autre téléphone", "Respondent — Other phone", ALL),
    # Step 3 — organisation
    ("ent_raison", "Raison sociale", "Company name", "ent"),
    ("ent_statut", "Statut juridique", "Legal status", "ent"),
    ("coop_nom", "Nom de la coopérative", "Cooperative name", "coop"),
    ("coop_type", "Type de coopérative", "Cooperative type", "coop"),
    ("ctd_type", "Type de CTD", "RLA type", "ctd"),
    ("ctd_nom", "Nom de la CTD", "RLA name", "ctd"),
    ("ong_nom", "Nom de l'ONG", "NGO name", "ong"),
    ("ong_enreg", "N° d'enregistrement (agrément)", "Registration no. (approval)", "ong"),
    ("adm_nom", "Nom de l'administration", "Administration name", "adm"),
    ("pp_type", "Type d'entité (projet, programme, structure sous-tutelle…)", "Type of entity", "pp"),
    ("pp_nom", "Nom du projet / programme / structure", "Project / programme / structure name", "pp"),
    ("cfp_nom", "Nom du CFP", "VTC name", "cfp"),
    ("cfp_type", "Type de CFP", "VTC type", "cfp"),
    ("cfp_ordre", "Ordre d'enseignement", "Education system", "cfp"),
    ("cfp_situation", "Situation du centre (fonctionnel ou non)", "Centre status (functional or not)", "cfp"),
    ("cfp_raison_nf", "Raison de non-fonctionnement", "Reason not functional", "cfp"),
    ("sigle", "Sigle", "Acronym", "adm pp cfp"),
    ("niu", "N° contribuable (NIU)", "Taxpayer no. (NIU)", "ent coop ctd ong cfp"),
    ("cnps", "N° CNPS (matricule employeur)", "CNPS no. (employer number)", "ent coop ctd ong cfp"),
    ("annee", "Année de création / d'ouverture", "Year established / opened", "coop ctd ong cfp"),
    ("activite", "Activité principale", "Main activity", "ent coop"),
    ("branche", "Branche d'activité", "Business branch", "ent coop"),
    ("act_sec", "Activité secondaire", "Secondary activity", "ent"),
    ("mission", "Mission principale", "Main mission", "ong adm pp"),
    ("adresse", "Adresse du siège / du CFP", "Office / VTC address", ALL),
    ("tel_etab", "Téléphone de l'établissement", "Establishment phone", ALL),
    ("tel_etab2", "Autre téléphone de l'établissement", "Other establishment phone", ALL),
    ("bp", "Boîte postale", "P.O. box", ALL),
    ("capital", "Capital social (XAF)", "Share capital (XAF)", "ent"),
    ("maison_mere", "Maison mère", "Parent company", "ent"),
    ("promo_nom", "Promoteur — Nom", "Promoter — Name", "cfp"),
    ("promo_sexe", "Promoteur — Sexe", "Promoter — Sex", "cfp"),
    ("promo_tel", "Promoteur — Téléphones", "Promoter — Phones", "cfp"),
    # Step 4 — location
    ("region", "Région", "Region", ALL),
    ("dept", "Département", "Division", ALL),
    ("arr", "Arrondissement", "Subdivision", ALL),
    ("milieu", "Milieu (urbain / rural)", "Area (urban / rural)", ALL),
    ("secteur", "Secteur d'activité", "Sector of activity", ALL),
    # Step 5
    ("password", "Mot de passe et confirmation", "Password and confirmation", ALL),
]
lst("field", FIELDS)
lst("field_none", [("none", "Aucun", "None", ALL)] + FIELDS)
OPTION_LISTS = [
    ("l_type_org", "Étape 1 — liste des types d'organisation", "Step 1 — types of organisation", ALL),
    ("l_fonction", "Fonction du répondant", "Respondent's position", ALL),
    ("l_statut", "Statut juridique", "Legal status", "ent"),
    ("l_coop", "Type de coopérative", "Cooperative type", "coop"),
    ("l_ctd", "Type de CTD", "RLA type", "ctd"),
    ("l_pp", "Type d'entité", "Type of entity", "pp"),
    ("l_cfp", "Type de CFP", "VTC type", "cfp"),
    ("l_ordre", "Ordre d'enseignement", "Education system", "cfp"),
    ("l_nf", "Raisons de non-fonctionnement", "Reasons not functional", "cfp"),
    ("l_geo", "Régions / départements / arrondissements", "Regions / divisions / subdivisions", ALL),
    ("l_secteur", "Secteurs d'activité", "Sectors of activity", ALL),
]
lst("optlist", OPTION_LISTS)

FILTER = "entities = 'all' or selected(entities, ${entity_type})"
FILTER_NONE = FILTER

# ---------------------------------------------------------------- survey
for meta in ("start", "end", "today"):
    survey.append({"type": meta, "name": meta})

note("intro",
     "**Évaluation de l'inscription autonome d'une entité sur CAMLEAP**\n\n"
     "Remplissez ce formulaire **une fois par inscription simulée**, juste après l'avoir terminée "
     "(ou abandonnée). Vous jouez le représentant d'une entité : jugez la plateforme comme le ferait "
     "ce représentant, puis donnez votre avis d'agent de l'ONEFOP.",
     "**Evaluation of an entity's self-registration on CAMLEAP**\n\n"
     "Fill in this form **once per simulated registration**, right after finishing (or abandoning) it. "
     "You play an entity's representative: judge the platform as that person would, then give your "
     "view as ONEFOP staff.")

# A — session
begin("a_session", "A. Le test", "A. The test")
q("text", "tester_name", "Votre nom (ou code testeur)", "Your name (or tester code)", req=True)
q("text", "tester_service", "Votre service / unité à l'ONEFOP", "Your ONEFOP department / unit", req=False)
q("select_one entity", "entity_type", "Quel type d'entité avez-vous simulé ?", "Which type of entity did you simulate?")
q("text", "entity_name", "Nom fictif de l'entité utilisé pour l'inscription",
  "Fictitious entity name used for the registration",
  hint=("Permet de retrouver le dossier sur la plateforme.", "Lets us find the file on the platform."), req=True)
q("text", "entity_email", "Email du répondant utilisé pour l'inscription", "Respondent e-mail used for the registration", req=False)
q("select_one browser", "browser", "Navigateur utilisé", "Browser used")
q("text", "browser_other", "Précisez le navigateur", "Specify the browser", req=True, relevant="${browser} = 'autre'")
q("select_one screen", "screen", "Taille de votre écran", "Your screen size")
q("select_one network", "network", "Connexion Internet utilisée", "Internet connection used")
q("select_one lang", "language", "Langue de l'interface pendant l'inscription", "Interface language during registration")
end("a_session")

# B — outcome
begin("b_outcome", "B. Déroulement et résultat", "B. How it went")
q("select_one outcome", "outcome", "Avez-vous pu aller jusqu'au bout de l'inscription ?", "Did you complete the registration?")
q("select_one step", "block_step", "À quelle étape avez-vous été bloqué(e) ou avez-vous abandonné ?",
  "At which step were you blocked or did you give up?", relevant="${outcome} != 'full'")
q("text", "block_desc", "Décrivez ce qui vous a bloqué(e)", "Describe what blocked you",
  relevant="${outcome} != 'full'", req=True, appearance="multiline")
q("text", "receipt_id", "Identifiant ou numéro affiché sur le récépissé", "ID or number shown on the receipt",
  relevant="${outcome} = 'full'", req=False)
q("integer", "duration", "Combien de minutes l'inscription vous a-t-elle pris (environ) ?",
  "About how many minutes did the registration take?", constraint=". >= 1 and . <= 240",
  cmsg=("Entre 1 et 240 minutes.", "Between 1 and 240 minutes."))
q("integer", "attempts", "Combien de tentatives avant de réussir (ou d'abandonner) ?",
  "How many attempts before succeeding (or giving up)?", constraint=". >= 1 and . <= 20",
  cmsg=("Entre 1 et 20.", "Between 1 and 20."))
q("select_one yesno", "needed_help", "Avez-vous eu besoin d'aide ?", "Did you need help?")
q("select_multiple help_src", "help_src", "Auprès de qui ou de quoi ?", "From whom or what?", relevant="${needed_help} = 'oui'")
q("select_multiple incident", "incidents", "Quels incidents avez-vous rencontrés ?", "Which incidents did you encounter?",
  constraint="not(selected(., 'none') and count-selected(.) > 1)",
  cmsg=("« Aucun incident » ne peut pas être combiné avec un autre choix.",
        "\"No incident\" cannot be combined with another choice."))
q("text", "incident_desc", "Décrivez les incidents (étape, champ, message affiché)",
  "Describe the incidents (step, field, message shown)",
  relevant="count-selected(${incidents}) > 0 and not(selected(${incidents}, 'none'))", req=True, appearance="multiline")
q("image", "incident_shot", "Capture d'écran de l'incident (facultatif)", "Screenshot of the incident (optional)", parameters="max-pixels=1600",
  relevant="count-selected(${incidents}) > 0 and not(selected(${incidents}, 'none'))", req=False)
q("select_one email_rx", "email_rx", "Avez-vous reçu un email de confirmation ?", "Did you receive a confirmation e-mail?",
  relevant="${outcome} = 'full'")
q("select_one seq", "seq", "Dans l'ensemble, créer ce compte a été…", "Overall, creating this account was…",
  hint=("1 = très difficile, 7 = très facile", "1 = very difficult, 7 = very easy"), appearance="likert")
end("b_outcome")

# C — step by step
begin("c_ease", "C1. Facilité de chaque étape", "C1. Ease of each step", appearance="table-list")
for code, fr, en in STEPS:
    q("select_one ease", f"ease_{code}", fr, en)
end("c_ease")
begin("c_clear", "C2. Les consignes et libellés de chaque étape étaient clairs",
      "C2. The instructions and labels of each step were clear", appearance="table-list")
for code, fr, en in STEPS:
    q("select_one agree", f"clear_{code}", fr, en)
end("c_clear")
begin("c_comments", "C3. Commentaires par étape (facultatif)", "C3. Comments per step (optional)")
note("c_comments_note", "Notez ce qui vous a gêné(e) ou plu, étape par étape. Laissez vide si rien à signaler.",
     "Note what bothered or pleased you, step by step. Leave blank if nothing to report.")
for code, fr, en in STEPS:
    q("text", f"comment_{code}", fr, en, req=False, appearance="multiline",
      relevant=f"${{ease_{code}}} != 'na'")
end("c_comments")

# D — content
begin("d_content", "D. Les informations demandées", "D. The information requested")
note("d_note", "Répondez en pensant à un **vrai représentant** de ce type d'entité. Seuls les champs que "
               "votre type d'entité a vus sont proposés.",
     "Answer thinking of a **real representative** of this type of entity. Only the fields your entity type saw are listed.")
q("select_one agree", "type_obvious", "À l'étape 1, un vrai représentant reconnaîtrait immédiatement son type d'organisation.",
  "At step 1, a real representative would immediately recognise their type of organisation.")
q("select_multiple entity", "type_confused", "Avec quel(s) autre(s) type(s) pourrait-il hésiter ?",
  "Which other type(s) might they hesitate with?", relevant="${type_obvious} = '1' or ${type_obvious} = '2' or ${type_obvious} = '3'",
  req=False)
q("select_one realism", "realism", "Un vrai représentant aurait-il toutes les informations demandées sous la main ?",
  "Would a real representative have all the requested information at hand?")
q("select_multiple field", "f_unavailable", "Lesquelles risque-t-il de ne pas avoir, ou de devoir chercher ?",
  "Which ones might they not have, or need to look up?", relevant="${realism} != 'yes'", choice_filter=FILTER)
q("select_multiple field_none", "f_unclear", "Quels libellés sont ambigus ou difficiles à comprendre ?",
  "Which labels are ambiguous or hard to understand?", choice_filter=FILTER,
  constraint="not(selected(., 'none') and count-selected(.) > 1)",
  cmsg=("« Aucun » ne peut pas être combiné avec un autre choix.", "\"None\" cannot be combined with another choice."))
q("select_multiple field_none", "f_useless", "Quelles informations vous semblent inutiles pour l'ONEFOP ou trop intrusives ?",
  "Which information seems useless to ONEFOP or too intrusive?", choice_filter=FILTER,
  constraint="not(selected(., 'none') and count-selected(.) > 1)",
  cmsg=("« Aucun » ne peut pas être combiné avec un autre choix.", "\"None\" cannot be combined with another choice."))
q("select_multiple field_none", "f_make_optional", "Quels champs obligatoires devraient être facultatifs ?",
  "Which required fields should be optional?", choice_filter=FILTER,
  constraint="not(selected(., 'none') and count-selected(.) > 1)",
  cmsg=("« Aucun » ne peut pas être combiné avec un autre choix.", "\"None\" cannot be combined with another choice."))
q("select_multiple field_none", "f_make_required", "Quels champs facultatifs devraient être obligatoires ?",
  "Which optional fields should be required?", choice_filter=FILTER,
  constraint="not(selected(., 'none') and count-selected(.) > 1)",
  cmsg=("« Aucun » ne peut pas être combiné avec un autre choix.", "\"None\" cannot be combined with another choice."))
q("text", "f_why", "Expliquez vos choix ci-dessus (champ et raison)", "Explain your choices above (field and reason)",
  req=False, appearance="multiline")
q("select_one yesno", "lists_ok", "Les listes de choix contenaient-elles toujours la bonne réponse pour cette entité ?",
  "Did the option lists always contain the right answer for this entity?")
q("select_multiple optlist", "lists_bad", "Quelle(s) liste(s) étai(en)t incomplète(s) ou inadaptée(s) ?",
  "Which list(s) were incomplete or unsuitable?", relevant="${lists_ok} = 'non'", choice_filter=FILTER)
q("text", "lists_missing", "Quelles options manquaient ou étaient mal formulées ?", "Which options were missing or badly worded?",
  relevant="${lists_ok} = 'non'", req=True, appearance="multiline")
q("select_one yesno", "info_missing_yn", "Manque-t-il une information dont l'ONEFOP a besoin pour identifier ou contacter cette entité ?",
  "Is any information ONEFOP needs to identify or contact this entity missing?")
q("text", "info_missing", "Laquelle, et pourquoi ?", "Which, and why?", relevant="${info_missing_yn} = 'oui'", req=True,
  appearance="multiline")
q("select_one score5", "content_score", "Note globale de la pertinence des informations demandées",
  "Overall score for the relevance of the information requested")
end("d_content")

# E — UX
UX = [
    ("ux_progress", "Je savais toujours à quelle étape j'étais et ce qui restait à faire.",
     "I always knew which step I was on and what was left."),
    ("ux_buttons", "Les boutons Continuer et Retour étaient faciles à trouver et se comportaient comme prévu.",
     "The Continue and Back buttons were easy to find and behaved as expected."),
    ("ux_required", "Je voyais clairement quels champs étaient obligatoires.", "I could clearly see which fields were required."),
    ("ux_errors", "Les messages d'erreur m'indiquaient précisément quoi corriger.", "Error messages told me exactly what to fix."),
    ("ux_format", "Les formats attendus (téléphone, NIU, année, email) étaient clairs avant de saisir.",
     "The expected formats (phone, NIU, year, e-mail) were clear before typing."),
    ("ux_readable", "Le texte était lisible (taille, contraste) sans zoomer.", "Text was readable (size, contrast) without zooming."),
    ("ux_layout", "La page était bien organisée, sans éléments superflus.", "The page was well organised, with nothing superfluous."),
    ("ux_keyboard", "Je pouvais remplir le formulaire au clavier (touche Tab) sans difficulté.",
     "I could fill in the form with the keyboard (Tab key) without difficulty."),
    ("ux_memory", "Mes réponses étaient conservées quand je revenais en arrière ou rechargeais la page.",
     "My answers were kept when I went back or reloaded the page."),
    ("ux_password", "Les règles du mot de passe étaient claires.", "The password rules were clear."),
    ("ux_review", "Le récapitulatif m'a permis de vérifier et corriger mes réponses facilement.",
     "The summary let me check and correct my answers easily."),
    ("ux_next", "Le récépissé final expliquait clairement la suite (validation par l'ONEFOP, connexion).",
     "The final receipt clearly explained what happens next (ONEFOP validation, sign-in)."),
    ("ux_lang", "La langue était cohérente du début à la fin.", "The language was consistent from start to finish."),
    ("ux_trust", "La plateforme inspire confiance, comme un service officiel de l'État.",
     "The platform inspires trust, like an official government service."),
]
begin("e_ux", "E. Interface et expérience utilisateur", "E. Interface and user experience", appearance="table-list")
for name, fr, en in UX:
    q("select_one agree", name, fr, en)
end("e_ux")

begin("f_perf", "F. Performance", "F. Performance")
q("select_one speed", "speed_pages", "Vitesse d'affichage des pages et des listes", "Speed of pages and lists")
q("select_one submit_wait", "submit_wait", "Temps d'attente après avoir cliqué sur le bouton final de soumission",
  "Waiting time after clicking the final submit button")
end("f_perf")

# G — SUS
SUS = [
    ("Je pense que j'utiliserais volontiers ce système régulièrement.", "I think that I would like to use this system frequently."),
    ("J'ai trouvé ce système inutilement complexe.", "I found the system unnecessarily complex."),
    ("J'ai trouvé ce système facile à utiliser.", "I thought the system was easy to use."),
    ("Je pense que j'aurais besoin de l'aide d'un technicien pour utiliser ce système.",
     "I think that I would need the support of a technical person to be able to use this system."),
    ("J'ai trouvé que les différentes fonctions de ce système étaient bien intégrées.",
     "I found the various functions in this system were well integrated."),
    ("J'ai trouvé qu'il y avait trop d'incohérences dans ce système.", "I thought there was too much inconsistency in this system."),
    ("J'imagine que la plupart des gens apprendraient très rapidement à utiliser ce système.",
     "I would imagine that most people would learn to use this system very quickly."),
    ("J'ai trouvé ce système très lourd à utiliser.", "I found the system very cumbersome to use."),
    ("Je me suis senti(e) très en confiance en utilisant ce système.", "I felt very confident using the system."),
    ("J'ai dû apprendre beaucoup de choses avant de pouvoir utiliser ce système.",
     "I needed to learn a lot of things before I could get going with this system."),
]
begin("g_sus", "G. Échelle d'utilisabilité (SUS) — en vous mettant à la place du représentant",
      "G. System Usability Scale (SUS) — as the representative", appearance="table-list")
for i, (fr, en) in enumerate(SUS, 1):
    q("select_one sus", f"sus{i}", fr, en)
end("g_sus")
odd = " + ".join(f"(number(${{sus{i}}}) - 1)" for i in (1, 3, 5, 7, 9))
even = " + ".join(f"(5 - number(${{sus{i}}}))" for i in (2, 4, 6, 8, 10))
q("calculate", "sus_score", req=False, calculation=f"({odd} + {even}) * 2.5")

# H — overall + proposals
begin("h_overall", "H. Appréciation globale", "H. Overall assessment")
note("sus_show", "Score SUS calculé : **${sus_score}** / 100", "Calculated SUS score: **${sus_score}** / 100")
q("select_one score5", "score_overall", "Note globale du parcours d'inscription", "Overall score for the registration journey")
q("select_one ready", "ready", "Selon vous, ce parcours est-il prêt pour de vraies entités lors du pilote ?",
  "In your view, is this journey ready for real entities in the pilot?")
q("text", "best", "Ce qui a le mieux fonctionné", "What worked best", req=False, appearance="multiline")
q("text", "worst", "Les 3 problèmes les plus importants, par ordre de gravité",
  "The 3 most important problems, in order of severity", req=True, appearance="multiline")
end("h_overall")

q("begin_repeat", "props", "I. Propositions d'amélioration", "I. Improvement proposals", req=False)
note("props_note", "Une fiche par proposition. Ajoutez-en autant que nécessaire.",
     "One entry per proposal. Add as many as needed.")
q("select_one step_general", "prop_step", "Où ?", "Where?")
q("select_one prop_cat", "prop_cat", "Type de proposition", "Type of proposal")
q("select_one field", "prop_field", "Champ concerné (facultatif)", "Field concerned (optional)", req=False,
  relevant="${prop_step} = 's2' or ${prop_step} = 's3' or ${prop_step} = 's4' or ${prop_step} = 's5'",
  choice_filter=FILTER)
q("text", "prop_problem", "Problème constaté", "Problem observed", req=True, appearance="multiline")
q("text", "prop_solution", "Amélioration proposée", "Proposed improvement", req=True, appearance="multiline")
q("select_one priority", "prop_priority", "Priorité", "Priority")
q("image", "prop_shot", "Capture d'écran (facultatif)", "Screenshot (optional)", req=False, parameters="max-pixels=1600")
survey.append({"type": "end_repeat", "name": "props"})

# ---------------------------------------------------------------- write
wb = Workbook()
ws = wb.active
ws.title = "survey"
ws.append(COLS)
for r in survey:
    ws.append([r.get(c, "") for c in COLS])
wc = wb.create_sheet("choices")
CH = ["list_name", "name", FR, EN, "entities"]
wc.append(CH)
for r in choices:
    wc.append([r.get(c, "") for c in CH])
wsx = wb.create_sheet("settings")
wsx.append(["form_title", "form_id", "version", "default_language", "style"])
wsx.append(["CAMLEAP — Évaluation de l'inscription autonome des entités", "camleap_eval_inscription_autonome",
            "2026102101", "Français (fr)", "pages"])

hdr = Font(name="Arial", bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="1F4E3D")
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font, cell.fill = hdr, fill
    sheet.freeze_panes = "A2"
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 22
for c in ("C", "D"):
    ws.column_dimensions[c].width = 55
    wc.column_dimensions[c].width = 45
wb.save(OUT)
print(f"survey rows: {len(survey)}, choices: {len(choices)}")
